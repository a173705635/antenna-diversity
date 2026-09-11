from pathlib import Path
import json,hashlib,re
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH,WD_BREAK,WD_TAB_ALIGNMENT
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[3]
WORK=Path(__file__).parent
DEST=ROOT/'GNSS天线分集可行性与验证综合报告.docx'
DATA=ROOT/'docs/reports/data'
checks=[]
for manifest in DATA.rglob('sha256.json'):
    hashes=json.loads(manifest.read_text(encoding='utf-8-sig'))
    if not isinstance(hashes,dict): continue
    count=0
    for name,digest in hashes.items():
        if not isinstance(digest,str) or len(digest)!=64: continue
        f=manifest.parent/name
        if not f.exists(): raise ValueError(str(f))
        assert hashlib.sha256(f.read_bytes()).hexdigest()==digest,str(f)
        count+=1
    checks.append({'manifest':str(manifest.relative_to(ROOT)),'verified':count})
(WORK/'data_verification.json').write_text(json.dumps(checks,indent=2),encoding='utf-8')

doc=Document()
sec=doc.sections[0]
sec.page_width=Cm(21);sec.page_height=Cm(29.7)
sec.top_margin=sec.bottom_margin=sec.right_margin=Cm(2.5);sec.left_margin=Cm(3)
sec.header_distance=sec.footer_distance=Cm(1.25)
def font(rpr):
    f=rpr.find(qn('w:rFonts'))
    if f is None: f=OxmlElement('w:rFonts');rpr.insert(0,f)
    f.attrib.clear()
    for slot,value in [('ascii','Times New Roman'),('hAnsi','Times New Roman'),('eastAsia','宋体'),('cs','Times New Roman')]: f.set(qn('w:'+slot),value)
def noindent(p):
    ind=p.get_or_add_pPr().find(qn('w:ind'))
    if ind is None: ind=OxmlElement('w:ind');p.get_or_add_pPr().append(ind)
    ind.set(qn('w:firstLineChars'),'0');ind.set(qn('w:firstLine'),'0')
for name,size in [('Normal',12),('Title',22),('Heading 1',18),('Heading 2',15),('Heading 3',14),('Caption',10.5),('Footer',10.5),('TOC Heading',18)]:
    st=doc.styles[name];font(st.element.get_or_add_rPr());st.font.size=Pt(size);st.font.color.rgb=RGBColor(0,0,0)
    pf=st.paragraph_format;pf.space_before=Pt(0);pf.space_after=Pt(0);pf.line_spacing=1.5
    if name=='Normal':
        pf.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
        ind=OxmlElement('w:ind');ind.set(qn('w:firstLineChars'),'200');ind.set(qn('w:firstLine'),'480');st.element.get_or_add_pPr().append(ind)
    else:
        noindent(st.element)
        pf.alignment=WD_ALIGN_PARAGRAPH.CENTER if name in ('Title','Heading 1','Caption','Footer','TOC Heading') else WD_ALIGN_PARAGRAPH.LEFT
    if name.startswith('Heading'):
        st.font.bold=True;pf.keep_with_next=True;pf.space_before=Pt(12);pf.space_after=Pt(6)
        if name=='Heading 1':pf.page_break_before=False
    if name in ('Title','TOC Heading'): st.font.bold=True;st.font.italic=False
toc_style=doc.styles['TOC Heading']
toc_style.base_style=doc.styles['Normal']
toc_style.paragraph_format.page_break_before=False
pp=toc_style.element.get_or_add_pPr()
np=OxmlElement('w:numPr');ni=OxmlElement('w:numId');ni.set(qn('w:val'),'0');np.append(ni);pp.append(np)
ol=OxmlElement('w:outlineLvl');ol.set(qn('w:val'),'9');pp.append(ol)
for st in doc.styles:
    for border in list(st.element.iter(qn('w:pBdr'))):border.getparent().remove(border)
for level in range(1,4):
    name=f'TOC {level}'
    if name not in doc.styles: doc.styles.add_style(name,1)
    st=doc.styles[name];font(st.element.get_or_add_rPr());st.font.size=Pt(12);noindent(st.element)
    st.paragraph_format.space_after=Pt(0);st.paragraph_format.line_spacing=1.5
defaults=doc.styles.element.find(qn('w:docDefaults'))
if defaults is not None:
    r=defaults.find(qn('w:rPrDefault')).find(qn('w:rPr'));font(r)
numbering=doc.part.numbering_part.element
abstract=OxmlElement('w:abstractNum');abstract.set(qn('w:abstractNumId'),'90')
for level in range(3):
    lv=OxmlElement('w:lvl');lv.set(qn('w:ilvl'),str(level))
    for tag,val in [('start','1'),('numFmt','decimal'),('pStyle',f'Heading{level+1}'),('lvlText','第%1章' if level==0 else '.'.join('%'+str(k+1) for k in range(level+1))),('suff','space')]:
        e=OxmlElement('w:'+tag);e.set(qn('w:val'),val);lv.append(e)
    abstract.append(lv)
numbering.append(abstract)
num=OxmlElement('w:num');num.set(qn('w:numId'),'90');e=OxmlElement('w:abstractNumId');e.set(qn('w:val'),'90');num.append(e);numbering.append(num)
for level in range(3):
    pp=doc.styles[f'Heading {level+1}'].element.get_or_add_pPr();np=OxmlElement('w:numPr')
    for tag,val in [('ilvl',str(level)),('numId','90')]:e=OxmlElement('w:'+tag);e.set(qn('w:val'),val);np.append(e)
    pp.append(np)
def field(p,code,cache=''):
    run=p.add_run();a=OxmlElement('w:fldChar');a.set(qn('w:fldCharType'),'begin');run._r.append(a)
    r=p.add_run();a=OxmlElement('w:instrText');a.set(qn('xml:space'),'preserve');a.text=' '+code+' ';r._r.append(a)
    r=p.add_run();a=OxmlElement('w:fldChar');a.set(qn('w:fldCharType'),'separate');r._r.append(a)
    p.add_run(cache)
    r=p.add_run();a=OxmlElement('w:fldChar');a.set(qn('w:fldCharType'),'end');r._r.append(a)
def para(text,style=None):return doc.add_paragraph(text,style)
def mt(text,plain=False):
    r=OxmlElement('m:r')
    if plain:
        pr=OxmlElement('m:rPr');sty=OxmlElement('m:sty');sty.set(qn('m:val'),'p');pr.append(sty);r.append(pr)
    t=OxmlElement('m:t');t.text=text;r.append(t);return r
def node(tag,*items):
    e=OxmlElement('m:'+tag)
    for i in items:
        if not isinstance(i,str) and i.tag==qn('m:e') and tag in ('num','den','e','sub','sup'):
            for child in list(i):e.append(child)
        else:e.append(mt(i) if isinstance(i,str) else i)
    return e
def frac(a,b):return node('f',node('num',a),node('den',b))
def sub(a,b):return node('sSub',node('e',a),node('sub',b))
def sup(a,b):return node('sSup',node('e',a),node('sup',b))
def seq(*items):return node('e',*items)
def summand(body,upper):
    n=OxmlElement('m:nary');pr=OxmlElement('m:naryPr')
    ch=OxmlElement('m:chr');ch.set(qn('m:val'),'∑');pr.append(ch);n.append(pr)
    n.append(node('sub','k=1'));n.append(node('sup',upper));n.append(node('e',body));return n
EQ={
 'wavelength':[mt('λ = '),frac('c','f')],
 'carrier':[mt('L(t) = −'),frac('λ','2π'),mt('Φ(t)')],
 'covariance':[sub('R','x'),mt(' = '),frac(seq('X',sup('X','H')),'N'),mt('， D = diag('),sup(sub('R','11'),'1/2'),mt('，…)， C = '),sup('D','−1'),sub('R','x'),sup('D','−1')],
 'blind':[sub('P','⊥'),mt(' = I − '),sub('v','J'),sup(sub('v','J'),'H'),mt('， w = '),frac(seq(sub('P','⊥'),sub('e','r')),seq(sup(sub('e','r'),'H'),sub('P','⊥'),sub('e','r')))],
 'mvdr':[mt('w = '),frac(seq(sup('R̃','−1'),'a'),seq(sup('a','H'),sup('R̃','−1'),'a')),mt('， R̃ = R + δI')],
 'gf':[mt('G = '),sub('SDCP',mt('L1',True)),mt(' − '),sub('SDCP',mt('L5',True)),mt('  （距离单位均为米）',True)],
 'sdcp':[mt('r = '),sub(mt('wrap',True),'λ'),mt('(q − g − b)， −'),frac('λ','2'),mt(' ≤ r < '),frac('λ','2')],
 'window':[sub('r̄','W'),mt(' = '),frac('λ','2π'),mt('arg',True),mt('['),frac('1','N'),summand(sup('e',seq('j',frac('2π','λ'),sub('r','k'))),'N'),mt(']')],
 'threshold':[sub('T','W'),mt(' = '),sub('Q','0.99'),mt('(|'),sub('r̄','W'),mt('|；独立LOS标定样本)',True)],
 'adaptive':[sub('T','i'),mt('(t) = '),sub('B','i'),mt('(t) + K'),sub('σ̂','i,W'),mt('(t)')],
 'fcw':[mt('ΔΦ = '),frac('2π',sup('2','30')),summand(seq(sub(mt('FCW',True),'k'),sub('N','clk,k')),'K')],
 'angle':[mt('ΔD(θ) ≈ d[sin θ − sin '),sub('θ','0'),mt(']')]
}
chapter=0;counts={'Table':0,'Equation':0};bid=100
def caption(p,kind):
    global bid
    counts[kind]+=1;bid+=1;name=f'{kind}_{chapter}_{counts[kind]}'
    start=OxmlElement('w:bookmarkStart');start.set(qn('w:id'),str(bid));start.set(qn('w:name'),name);p._p.append(start)
    field(p,'STYLEREF 1 \\s',str(chapter));p.add_run('-');field(p,f'SEQ {kind} \\* ARABIC \\s 1',str(counts[kind]))
    end=OxmlElement('w:bookmarkEnd');end.set(qn('w:id'),str(bid));p._p.append(end)
    return name
def equation(key):
    p=para('');noindent(p._p);p.paragraph_format.tab_stops.add_tab_stop(Cm(7.75),WD_TAB_ALIGNMENT.CENTER);p.paragraph_format.tab_stops.add_tab_stop(Cm(15.5),WD_TAB_ALIGNMENT.RIGHT)
    p.add_run('\t');math=OxmlElement('m:oMath')
    for item in EQ[key]:math.append(item)
    p._p.append(math);p.add_run('\t（');caption(p,'Equation');p.add_run('）')
def table(headers,rows,title):
    p=para('相关项目与数值见表');ref=para('', 'Caption');ref.add_run('表 ');mark=caption(ref,'Table');ref.add_run(' '+title);ref.paragraph_format.keep_with_next=True
    field(p,f'REF {mark} \\h');p.add_run('。')
    t=doc.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    width=15.5/len(headers)
    for c in t.columns:c.width=Cm(width)
    pr=t._tbl.tblPr
    borders=OxmlElement('w:tblBorders')
    for tag in ('top','bottom','left','right','insideH','insideV'):
        x=OxmlElement('w:'+tag);x.set(qn('w:val'),'single' if tag in ('top','bottom') else 'nil');x.set(qn('w:sz'),'8');borders.append(x)
    pr.append(borders)
    for vals,row in [(headers,t.rows[0])]+[(r,t.add_row()) for r in rows]:
        trpr=row._tr.get_or_add_trPr();trpr.append(OxmlElement('w:cantSplit'))
        for c,text in zip(row.cells,vals):
            c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER;p=c.paragraphs[0];noindent(p._p)
            p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.line_spacing=1.15;p.paragraph_format.space_before=Pt(3);p.paragraph_format.space_after=Pt(3)
            rr=p.add_run(str(text));rr.font.size=Pt(10.5);font(rr._r.get_or_add_rPr())
            if row._tr is t.rows[0]._tr:rr.bold=True
    h=t.rows[0]._tr.get_or_add_trPr();h.append(OxmlElement('w:tblHeader'))
    for ri,row in enumerate(t.rows):
        if ri==0 or (len(t.rows)<=7 and ri<len(t.rows)-1):
            for cell in row.cells:
                for p in cell.paragraphs:p.paragraph_format.keep_with_next=True
    for c in t.rows[0].cells:
        b=OxmlElement('w:tcBorders');e=OxmlElement('w:bottom');e.set(qn('w:val'),'single');e.set(qn('w:sz'),'6');b.append(e);c._tc.get_or_add_tcPr().append(b)
    return t

title=para('GNSS天线分集可行性与验证综合报告','Title');title.paragraph_format.space_before=Pt(48);title.paragraph_format.space_after=Pt(24)
p=para('接收链路  多径检测  定位评估  工程自检','Subtitle');font(p.style.element.get_or_add_rPr());p.style.font.color.rgb=RGBColor(0,0,0);p.style.font.italic=False;p.alignment=WD_ALIGN_PARAGRAPH.CENTER;noindent(p._p)
p=para('2026年9月11日');p.alignment=WD_ALIGN_PARAGRAPH.CENTER;noindent(p._p);p.paragraph_format.space_after=Pt(30)
p=para('内容摘要');p.runs[0].bold=True;noindent(p._p)
para('本报告围绕GNSS多天线系统的观测可信度、差分多径检测及定位收益，整合架构与算法资料、已完成仿真实验和后续硬件验证方案。当前证据支持在已知几何与通道校准条件下检测强非共模多径，但弱多径、共模、长期偏差及标定迁移仍未解决。窗口平均能够降低随机波动，不能替代误差建模；定位收益也须与误剔除和几何损失共同评价。')
para('报告按理论、实现、实验设计、结果与边界展开，区分历史基线、当前验证和待实现功能。抗干扰作为独立技术路线说明，不与关闭外部干扰的多径专项混用性能结论。')
doc.add_page_break();para('目录','TOC Heading');field(para(''),'TOC \\o "1-3" \\h \\z \\u')
doc.add_page_break()
content=(WORK/'report_content.md').read_text(encoding='utf-8').splitlines()
i=0;tn=0
while i<len(content):
    line=content[i]
    if line.startswith('# '):
        chapter+=1;counts={'Table':0,'Equation':0};doc.add_heading(line[2:],1)
    elif line.startswith('## '):doc.add_heading(line[3:],2)
    elif line.startswith('@EQ '):equation(line.split()[1])
    elif line=='@SHORT_TABLE':
        rows=json.loads((DATA/'2026-09-11-sdcp-windows/summary.json').read_text())
        vals=[]
        for r in rows:
            weak=[x['late_prn1_fraction']*100 for x in r['cases'] if x['kind']=='weak'];fa=[x['late_all_false_fraction']*100 for x in r['cases'] if x['kind']=='los']
            vals.append([str(r['window_ms']),f"{r['threshold_mm']:.3f}",f"{r['los_mean_temporal_std_mm']:.3f}",'/'.join(f'{x:.1f}%' for x in weak),'/'.join(f'{x:.2f}%' for x in fa)])
        table(['窗长 ms','阈值 mm','时间STD mm','弱检出 两组','LOS误报 两组'],vals,'1秒记录的窗口对比')
    elif line=='@LONG_TABLE':
        rows=json.loads((DATA/'2026-09-11-sdcp-long/summary.json').read_text());vals=[]
        for w in (200,500,1000):
            subset=[r for r in rows if r['window_ms']==w];fa=[r['all_fraction']*100 for r in subset if r['kind']=='los']
            assert all(r['prn1_fraction']==(1. if r['kind'] in ('strong','static') else 0.) for r in subset)
            vals.append([w,f"{subset[0]['threshold_mm']:.3f}",'100% / 100%','0% / 0%','/'.join(f'{v:.2f}%' for v in fa)])
        table(['窗长 ms','阈值 mm','强检出 两组','弱检出 两组','LOS误报 两组'],vals,'3秒记录的后期窗口对比')
    elif line.startswith('|'):
        lines=[]
        while i<len(content) and content[i].startswith('|'):
            lines.append([v.strip() for v in content[i].strip('|').split('|')]);i+=1
        tn+=1;table(lines[0],lines[1:],['接收架构与能力','33 MHz采样下的定位结果','双频定位模式对比','SDCP软降权的定位收益与代价'][tn-1]);continue
    elif line.strip():para(line)
    i+=1

chapter+=1;counts={'Table':0,'Equation':0};doc.add_heading('来源索引与一致性核验',1)
doc.add_heading('材料整合与去重',2)
para('共纳入17份有效DOCX，按SHA-256识别出3组完全相同副本，合计14份独立内容。锁文件不作为正文材料。下表说明各来源的用途；同一内容不重复拼接，原文全部保留。')
inventory=json.loads((WORK/'inventory.json').read_text(encoding='utf-8'))
table(['来源文件','整合用途'],[[r['path'].replace('\\',' / '), '架构与历史方法' if r['id'] in (8,9,10,11,13) else '算法与工程验证依据'] for r in inventory],'DOCX材料清单')
doc.add_heading('需要统一或修正的表述',2)
for text in [
 'L5整周项统一乘L5波长。原跨频稳定性方案某式中的L1波长为符号笔误，不作为新模型沿用。',
 '最终FCW按实际生效时钟累计；更新次数固定不意味着可以用最后一个FCW代表整段。距离符号按底层擦除约定验证，不混用不同材料的正负号。',
 '参考角差分用于连续几何变化验证，不用于证明启动无多径；去参考常数只能保留变化，不能宣称完整静态检测能力。',
 '四天线六条两两差分仅有三个独立自由度，纯代数闭环不能独立识别多径。原五天线论文、四天线扩展和双天线实现分开陈述。',
 '旧报告中的680项、67项等测试数属于历史版本，不当作当前测试总数。理论角度正弦曲线、理想零陷及数值示例不当作硬件实测数据。',
 '未对专利授权日期、法律状态或覆盖范围作结论；使用A1公开文本的技术思想及仓库存档，不将专利说明书性能写成项目复现结果。']:
    para(text)
doc.add_heading('归档与复现入口',2)
para('本次文档生成前已校验现有SHA-256清单。校验保证归档字节一致，不证明模型物理正确。窗口表由JSON汇总直接生成；其他数值按对应专题报告与归档核对。分析脚本输出与历史实验各有配置，复现不得用新默认值覆盖旧条件。')
table(['研究内容','主要入口','结果归档目录后缀'],[
 ['相位基准','audit_phase_origin.py','phase-origin'],['SDCP定位','audit_sdcp_positioning.py','sdcp-positioning'],['校准取余','audit_calibrated_sdcp.py','calibrated-sdcp'],['短窗口','audit_sdcp_windows.py','sdcp-windows'],['3秒长窗','audit_sdcp_long.py','sdcp-long'],['双频定位','run_gps_l1_l5.py','dual-frequency-optimization']], '代码与归档索引')
para('代码位于python_sim_simplified；上表归档均位于docs/reports/data/2026-09-11-对应后缀。可运行各目录verify.py核查其提供的统计与哈希。本次为文档整合，没有重跑全部RF仿真，也没有改变检测或定位算法。')
doc.add_heading('参考资料',2)
refs=[
 'GNSS天线分集可行性研究报告_初稿.docx。根目录与docs/reports副本相同。',
 '4.3_多天线GNSS工程实现关键问题与约束.docx。项目工程分析材料。',
 'Blind_Nulling盲零陷原理章节.docx；docs/reports/Blind_Nulling实现方法_精简版.docx。',
 '5.2_MVDR波束形成算法_可行性报告章节.docx；docs/reports/MVDR实现方法_精简版.docx。',
 '6.1_Trimble多径检测方案_最终修订版.docx。项目专利技术整理。',
 '6.2_基于阵列载波相位姿态一致性的多径检测方案.docx。含论文整理与工程扩展。',
 '验证方案/基于最终FCW的载波NCO累加相位重建方案.docx。待硬件验证。',
 '验证方案/手机双天线跨频GNSS相位稳定性验证方案.docx。待硬件验证。',
 '验证方案/双天线跨频GNSS物理路径差角度差分验证方案.docx。待硬件验证。',
 'docs/reports/GNSS天线分集与空域抗干扰仿真_项目说明与基础理论_2026-08-29.docx。历史版本。',
 'docs/reports/GNSS多天线_BlindNull到MVDR_原理与方案整理.docx。',
 'docs/reports/2026-09-10-multipath-stage-summary.md及其引用的门限与边界报告；GNSS多径天线分集算法原理与实现说明.docx为早期实现背景。',
 'docs/reports/2026-09-11-dual-tracking-stability.md及2026-09-10-dual-tracking-phase.md。',
 'docs/reports/2026-09-11-gps-spp-validation.md及2026-09-11-code-carrier-float.md。',
 'docs/reports/2026-09-11-pseudorange-bias-diagnosis.md。',
 'docs/reports/2026-09-11-sampling-33mhz.md。',
 'docs/reports/2026-09-11-gps-l5.md及2026-09-11-dual-frequency-optimization.md。',
 'docs/reports/2026-09-11-phase-origin.md。',
 'docs/reports/2026-09-11-calibrated-sdcp.md。',
 'docs/reports/2026-09-11-sdcp-positioning.md。',
 'docs/reports/2026-09-11-sdcp-windows.md。',
 'docs/reports/2026-09-11-sdcp-long.md。',
 'docs/reports/2026-09-11-sdcp-self-check-feasibility.md。待实现设计。',
 'docs/reports/2026-09-11-trimble-single-frequency-validation.md。旧采样率单星SDCP验证。',
 'Clare A, Weisenburger S, Wallace G. Multipath Mitigation for GNSS Location Determination Using Dual Antennas. US 2024/0219581 A1, 2024-07-04. https://patents.google.com/patent/US20240219581A1/en 。另有项目存档US12656504B2.pdf。',
 'Ahmed M, Ballal T, Liu X, Al-Naffouri T Y. Multipath Detection and Mitigation from GNSS Observations Using Antenna Arrays. IEEE/ION PLANS, 2023. DOI: 10.1109/PLANS53410.2023.10139949。',
 'Ahmed A, Ballal T, Liu X, Ahmed M, Al-Naffouri T Y. GNSS Array-Based Multipath Detection Employing UKF on Manifolds. IEEE/ION PLANS, 2025: 610–618. DOI: 10.1109/PLANS61210.2025.11028455。']
refs.extend(['docs/reports/2026-09-08-multipath-baseline-audit.md。直达径基线与连续本地码诊断。','docs/reports/2026-09-08-continuous-source-oracle.md。连续信号源与理想参考。','docs/reports/2026-09-08-tracking-initialization.md。数据驱动初始化控制试验。','docs/reports/2026-09-08-detector-versus-oracle.md。时延字段与非网格方向反例。','docs/reports/2026-09-10-detection-boundaries.md。冻结门限与单因素边界。'])
refs.append('docs/reports/2026-09-08-detection-threshold-calibration.md。记录级顺序统计门限与独立评估区间。')
for n,ref in enumerate(refs,1):
    p=para(f'[{n}] {ref}');noindent(p._p);p.paragraph_format.line_spacing=1.15
field(sec.footer.paragraphs[0],'PAGE');noindent(sec.footer.paragraphs[0]._p)
setting=OxmlElement('w:updateFields');setting.set(qn('w:val'),'true');doc.settings.element.append(setting)
doc.core_properties.title='GNSS天线分集可行性与验证综合报告';doc.core_properties.subject='多径检测与接收链路验证';doc.core_properties.author=''
doc.save(DEST)
print(str(DEST));print('paragraphs',len(doc.paragraphs),'tables',len(doc.tables),'verified',sum(r['verified'] for r in checks))
