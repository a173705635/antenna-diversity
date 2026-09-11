from pathlib import Path
import sys,json
P=Path(__file__).parent; R=P.parents[2]
sys.path.append(str(P/'vendor'))
from docx import Document
from docxcompose.composer import Composer
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
main=R/'GNSS天线分集可行性与验证综合报告.docx'
d=Document(main);c=Composer(d)
for item in json.loads((P/'full_merge_manifest.json').read_text(encoding='utf-8')):
    if item['name']==main.name:continue
    d.add_page_break()
    p=d.add_paragraph('原文完整收录：'+item['name']);p.runs[0].bold=True
    d.add_paragraph('以下保留来源全文及原始表述。历史状态、理论示例和待验证内容请结合前部综合报告判读。')
    source=Document(R/item['name'])
    # Source headings already contain their chapter numbers as literal text.
    # Override numbering inherited from the destination's shared heading styles.
    for p in source.paragraphs:
        if p.style and p.style.name.lower().startswith('heading'):
            props=p._p.get_or_add_pPr()
            old=props.find(qn('w:numPr'))
            if old is not None: props.remove(old)
            num=OxmlElement('w:numPr'); nid=OxmlElement('w:numId')
            nid.set(qn('w:val'),'0'); num.append(nid); props.append(num)
    c.append(source)
c.save(R/'GNSS天线分集全量整合报告.docx')
script=(P/'update_word.ps1').read_text(encoding='utf-8-sig').replace('GNSS天线分集可行性与验证综合报告.docx','GNSS天线分集全量整合报告.docx').replace('report.pdf','full_merge.pdf')
# Do not regenerate source TOCs or renumber original source fields as a new report.
script=script.replace('$reportDoc.Fields.Update() | Out-Null','').replace('foreach ($toc in $reportDoc.TablesOfContents) { $toc.Update() }','')
(P/'render_full.ps1').write_text(script,encoding='utf-8-sig')
print('Composed all seven sources')
